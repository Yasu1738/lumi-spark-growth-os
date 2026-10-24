"""CSV・Excelの非破壊インポート。

- 元ファイルは読み取り専用で開き、取込前後のSHA-256が一致することを確認する。
- 既存の記事IDは上書きしない（スキップして報告する）。
- 空欄は None（欠損）として取り込み、0 に置き換えない。
- 旧台帳（note記事運用台帳 Ver.0.1「記事台帳」38列）の未対応列は legacy_json に保全する。
"""

import csv
import hashlib
import json
from pathlib import Path

from . import repository
from .validation import ValidationError, clean_text, is_blank, parse_datetime

# note記事運用台帳 Ver.0.1「記事台帳」シートのヘッダー（A〜AL列、38列）
LEGACY_LEDGER_COLUMNS = [
    "Article_ID", "公開日時", "タイトル", "URL", "主目的", "副目的", "テーマ", "想定読者",
    "感情温度", "冒頭タイプ", "結末タイプ", "核となる一文", "公開前反応予測点(25)",
    "24h PV", "24h スキ", "24h コメント", "24h スキ率",
    "72h PV", "72h スキ", "72h コメント", "72h スキ率",
    "フォロー増(参考)", "外部リンク反応", "問い合わせ・相談", "24h期限", "72h期限",
    "レビュー状態", "基準PV", "基準スキ率", "PV判定", "反応判定", "4象限", "総合判定",
    "残す要素", "変える要素", "次の仮説", "次アクション", "メモ",
]
LEGACY_TO_ARTICLE = {
    "Article_ID": "article_id", "公開日時": "published_at", "タイトル": "title", "URL": "url",
    "主目的": "primary_purpose", "副目的": "secondary_purpose", "テーマ": "theme",
    "想定読者": "target_reader", "感情温度": "emotion_temp", "核となる一文": "core_claim",
    "次の仮説": "pre_hypothesis", "メモ": "notes",
}
# 日次アカウントシート（16列）
LEGACY_ACCOUNT_COLUMNS = [
    "集計日時", "対象期間開始", "対象期間終了", "28日インプレッション", "28日PV", "28日スキ",
    "28日コメント", "28日フォロー増(正確値)", "28日投稿数", "前回窓差インプレッション", "前回窓差PV",
    "前回窓差スキ", "前回窓差コメント", "前回窓差フォロー増", "前回窓差投稿数", "メモ",
]

NATIVE_ARTICLE_COLUMNS = repository.ARTICLE_FIELDS[:-1]  # legacy_json は除く
NATIVE_METRIC_COLUMNS = ["article_id", "measured_at", "window_type", *repository.METRIC_FIELDS,
                         "source", "condition_note"]
NATIVE_LINK_COLUMNS = ["source_article_id", "target_article_id", "reason", "reader_value",
                       "anchor_text", "insert_position", "evidence", "hub_id"]


def file_sha256(path) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(65536), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _cell(value):
    if is_blank(value):
        return None
    return value.strip() if isinstance(value, str) else value


def read_table(path, sheet=None):
    """(ヘッダー, 行dictのリスト, シート名) を返す。元ファイルは書き換えない。"""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".csv":
        raw = path.read_bytes()
        for encoding in ("utf-8-sig", "cp932"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ValidationError("CSVの文字コードを判別できません（UTF-8またはShift_JISで保存してください）")
        reader = csv.reader(text.splitlines())
        rows = list(reader)
        sheet_name = None
    elif suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            if sheet is None:
                sheet = next((s for s in ("記事台帳", "日次アカウント") if s in wb.sheetnames), wb.sheetnames[0])
            if sheet not in wb.sheetnames:
                raise ValidationError(f"シート「{sheet}」がありません（{', '.join(wb.sheetnames)}）")
            rows = [list(r) for r in wb[sheet].iter_rows(values_only=True)]
        finally:
            wb.close()
        sheet_name = sheet
    else:
        raise ValidationError("対応形式は .csv / .xlsx / .xlsm です")
    if not rows:
        return [], [], sheet_name
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    records = []
    for row in rows[1:]:
        record = {h: _cell(row[i]) if i < len(row) else None for i, h in enumerate(headers) if h}
        if all(v is None for v in record.values()):
            continue  # テンプレートの空行
        records.append(record)
    return headers, records, sheet_name


def list_sheets(path):
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def detect_format(headers) -> str:
    hs = set(headers)
    if {"Article_ID", "タイトル", "24h PV"} <= hs:
        return "legacy_ledger"
    if {"集計日時", "28日PV"} <= hs:
        return "legacy_account"
    if {"source_article_id", "target_article_id"} <= hs:
        return "native_links"
    if {"article_id", "measured_at", "window_type"} <= hs:
        return "native_metrics"
    if {"article_id", "title"} <= hs:
        return "native_articles"
    return "unknown"


def _legacy_metric_rows(record):
    """旧台帳の24h/72h列を計測値として取り出す。値のない区分は作らない。"""
    out = []
    for window, pv_col, like_col, comment_col, deadline_col in (
        ("24h", "24h PV", "24h スキ", "24h コメント", "24h期限"),
        ("72h", "72h PV", "72h スキ", "72h コメント", "72h期限"),
    ):
        values = {
            "pv": record.get(pv_col), "likes": record.get(like_col), "comments": record.get(comment_col),
        }
        if all(is_blank(v) for v in values.values()):
            continue
        measured = record.get(deadline_col)
        note = f"旧台帳「{deadline_col}」列を計測日時として転記（実計測時刻と異なる可能性）"
        out.append((window, measured, values, note))
    return out


SAMPLE_PREFIX = "【サンプル】"


def import_file(conn, path, actor, sheet=None, dry_run=False, display_name=None, as_sample=False) -> dict:
    """ファイルを取り込み、レポートを返す。

    dry_run=True なら DB を変更しない。as_sample=True なら記事をサンプル（実績ではない）として登録する。
    タイトルが「【サンプル】」で始まる記事は、指定がなくてもサンプルとして登録する。
    """
    actor = repository._require_actor(actor)
    path = Path(path)
    sha_before = file_sha256(path)
    headers, records, sheet_name = read_table(path, sheet)
    fmt = detect_format(headers)
    report = {
        "file": display_name or path.name, "sheet": sheet_name, "format": fmt, "dry_run": dry_run,
        "rows_read": len(records), "inserted": 0, "skipped": [], "errors": [], "warnings": [],
        "metrics_inserted": 0, "sha256_before": sha_before, "as_sample": as_sample,
    }
    if fmt == "legacy_ledger":
        missing = [c for c in LEGACY_LEDGER_COLUMNS if c not in headers]
        extra = [h for h in headers if h and h not in LEGACY_LEDGER_COLUMNS]
        if missing:
            report["warnings"].append(f"旧台帳の想定列のうち不足: {', '.join(missing)}")
        if extra:
            report["warnings"].append(f"想定外の列（legacy_jsonに保全）: {', '.join(extra)}")

    conn.defer_commit = getattr(conn, "defer_commit", 0) + 1
    conn.execute("SAVEPOINT import_all")
    try:
        handler = {
            "legacy_ledger": _import_legacy_ledger,
            "legacy_account": _import_legacy_account,
            "native_articles": _import_native_articles,
            "native_metrics": _import_native_metrics,
            "native_links": _import_native_links,
        }.get(fmt)
        if handler is None:
            report["errors"].append("列構成を判別できません。README の「取込フォーマット」を確認してください")
        else:
            for index, record in enumerate(records, start=2):
                conn.execute("SAVEPOINT import_row")
                try:
                    handler(conn, record, actor, report, index)
                    conn.execute("RELEASE import_row")
                except ValidationError as exc:
                    conn.execute("ROLLBACK TO import_row")
                    conn.execute("RELEASE import_row")
                    report["errors"].append(f"{index}行目: {exc}")
        if dry_run:
            conn.execute("ROLLBACK TO import_all")
        conn.execute("RELEASE import_all")
    except Exception:
        conn.execute("ROLLBACK TO import_all")
        conn.execute("RELEASE import_all")
        raise
    finally:
        conn.defer_commit -= 1

    sha_after = file_sha256(path)
    report["sha256_after"] = sha_after
    report["source_unchanged"] = sha_before == sha_after
    conn.execute(
        "INSERT INTO import_batches(imported_at, actor, file_name, file_sha256, detected_format, dry_run,"
        " rows_read, rows_inserted, rows_skipped, report_json) VALUES (?,?,?,?,?,?,?,?,?,?)",
        (repository.now_iso(), actor, report["file"], sha_before, fmt, int(dry_run), report["rows_read"],
         report["inserted"], len(report["skipped"]), json.dumps(report, ensure_ascii=False, default=str)),
    )
    conn.commit()
    return report


def _next_legacy_id(conn):
    row = conn.execute(
        "SELECT article_id FROM articles WHERE article_id LIKE 'LEG-%' ORDER BY article_id DESC LIMIT 1"
    ).fetchone()
    number = int(row[0].split("-")[1]) + 1 if row else 1
    return f"LEG-{number:04d}"


def _insert_article(conn, data, actor, report, index):
    title = clean_text(data.get("title")) or ""
    if report["as_sample"] or title.startswith(SAMPLE_PREFIX):
        data["data_kind"] = "sample"
    article_id = clean_text(data.get("article_id"))
    if article_id and repository.get_article(conn, article_id):
        report["skipped"].append(f"{index}行目: 記事ID {article_id} は既存のため上書きしません")
        return None
    article = repository.create_article(conn, data, actor, strict_url=False)
    if article["url_status"] == "不正":
        report["warnings"].append(f"{index}行目: 不正URL「{article['url']}」を検出（URL状態=不正として登録）")
    report["inserted"] += 1
    return article


def _import_legacy_ledger(conn, record, actor, report, index):
    if is_blank(record.get("タイトル")) and is_blank(record.get("Article_ID")):
        return
    data = {field: record.get(col) for col, field in LEGACY_TO_ARTICLE.items()}
    if is_blank(data.get("article_id")):
        data["article_id"] = _next_legacy_id(conn)
        report["warnings"].append(f"{index}行目: Article_IDが空欄のため {data['article_id']} を採番")
    if is_blank(data.get("title")):
        data["title"] = "（タイトル未入力）"
        data["title_status"] = "未確認"
    else:
        data["title_status"] = "確認済み"
    # テーマ・主目的・感情温度が定義外の値なら、推測で変換せず legacy_json に残す
    from . import codes
    for field, allowed in (("theme", codes.THEMES), ("primary_purpose", codes.PURPOSES),
                           ("emotion_temp", codes.EMOTION_TEMPS)):
        if data.get(field) is not None and data[field] not in allowed:
            report["warnings"].append(f"{index}行目: {field}「{data[field]}」は定義外のため未設定（原文はlegacy_jsonに保全）")
            data[field] = None
    data["url_status"] = "未確認"
    data["source"] = f"旧台帳インポート: {report['file']}"
    data["verification_status"] = "一部確認"
    data["legacy_json"] = json.dumps(
        {k: v for k, v in record.items() if v is not None}, ensure_ascii=False, default=str)
    article = _insert_article(conn, data, actor, report, index)
    if article is None:
        return
    for window, measured, values, note in _legacy_metric_rows(record):
        if is_blank(measured):
            report["warnings"].append(f"{index}行目: {window}値がありますが計測日時（{window}期限）が空欄のため実績に取り込まず保全のみ")
            continue
        try:
            repository.add_metric(conn, article["article_id"], measured, window, values,
                                  f"旧台帳インポート: {report['file']}", actor, condition_note=note)
            report["metrics_inserted"] += 1
        except ValidationError as exc:
            report["warnings"].append(f"{index}行目: {window}値を実績に取り込めません（{exc}）。原文はlegacy_jsonに保全")


def _import_legacy_account(conn, record, actor, report, index):
    if is_blank(record.get("集計日時")):
        return
    follows = record.get("28日フォロー増(正確値)")
    data = {
        "measured_at": record.get("集計日時"), "period_start": record.get("対象期間開始"),
        "period_end": record.get("対象期間終了"), "imp": record.get("28日インプレッション"),
        "pv": record.get("28日PV"), "likes": record.get("28日スキ"), "comments": record.get("28日コメント"),
        "follows": follows, "posts": record.get("28日投稿数"), "period_note": "28日窓",
        "source": f"旧台帳インポート: {report['file']}", "notes": record.get("メモ"),
    }
    measured = parse_datetime(data["measured_at"], "集計日時")
    dup = conn.execute("SELECT 1 FROM account_snapshots WHERE measured_at = ? AND data_kind='real'",
                       (measured,)).fetchone()
    if dup:
        report["skipped"].append(f"{index}行目: 集計日時 {measured} は既存のため上書きしません")
        return
    repository.add_account_snapshot(conn, data, actor)
    report["inserted"] += 1


def _import_native_articles(conn, record, actor, report, index):
    data = {k: record.get(k) for k in NATIVE_ARTICLE_COLUMNS if k in record}
    data.setdefault("source", f"インポート: {report['file']}")
    _insert_article(conn, data, actor, report, index)


def _import_native_metrics(conn, record, actor, report, index):
    values = {k: record.get(k) for k in repository.METRIC_FIELDS}
    repository.add_metric(conn, record.get("article_id"), record.get("measured_at"),
                          clean_text(record.get("window_type")), values,
                          record.get("source") or f"インポート: {report['file']}", actor,
                          condition_note=record.get("condition_note"))
    report["inserted"] += 1
    report["metrics_inserted"] += 1


def _import_native_links(conn, record, actor, report, index):
    src, tgt = record.get("source_article_id"), record.get("target_article_id")
    issues = repository.check_link_candidate(conn, src, tgt)
    for issue in issues:
        if issue.startswith("不正URL"):
            report["warnings"].append(f"{index}行目: {issue}")
    blocking = [i for i in issues if not i.startswith("不正URL")]
    if blocking:
        report["skipped"].append(f"{index}行目: " + " / ".join(blocking))
        return
    repository.create_link(
        conn, src, tgt, actor, reason=record.get("reason"), reader_value=record.get("reader_value"),
        anchor_text=record.get("anchor_text"), insert_position=record.get("insert_position"),
        evidence=record.get("evidence"), hub_id=record.get("hub_id"),
        proposed_by=f"インポート: {report['file']}",
    )
    report["inserted"] += 1
